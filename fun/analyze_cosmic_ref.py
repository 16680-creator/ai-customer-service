# -*- coding: utf-8 -*-
"""Analyse 附件3 COSMIC split conventions, resolving merged-cell ranges properly."""
import sys
from collections import Counter, OrderedDict

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

SRC = (r"d:\Software\xwechat_files\wxid_8mnhlxo6ia3822_2fff\temp\RWTemp\2026-09"
       r"\33d7b6535ac286b8e16798e602ffa3b0\附件3：湖南个人安全-COSMIC功能点拆分表.xlsx")

VERBS = """查询 获取 生成 新增 保存 存管 归档 冻结 申请 提交 修改 更新 删除 汇总 推送 采集 处理
打印 上传 下载 下发 设置 计算 分析 通知 统计 展示 同步 迁移 预览 审批 对比 比对 导入 导出""".split()

wb = openpyxl.load_workbook(SRC, data_only=True)
print("sheets:", wb.sheetnames)

ws = wb["个人客户信息安全"]

# Map every covered cell of a merged range to its top-left value.
fill = {}
for rng in ws.merged_cells.ranges:
    tl = ws.cell(row=rng.min_row, column=rng.min_col).value
    for row in ws[rng.coord]:
        for cell in row:
            fill[(cell.row, cell.column)] = tl

header = {}
for c in ws[1]:
    if c.value:
        header[str(c.value).strip()] = c.column


def val(row, col_name):
    col = header[col_name]
    v = fill.get((row, col), ws.cell(row=row, column=col).value)
    return str(v).strip() if v is not None else ""


records = []
for r in range(2, ws.max_row + 1):
    dm = val(r, "数据移动类型")
    fp = val(r, "功能过程")
    if not dm and not fp:
        continue
    rec = {name: val(r, name) for name in
           ("一级模块", "二级模块", "三级模块", "功能过程", "子过程描述", "数据移动类型", "数据组")}
    records.append(rec)

print("数据移动(子过程)总行数:", len(records))
print("数据移动类型分布:", Counter(x["数据移动类型"] for x in records).most_common())

groups = OrderedDict()
for x in records:
    key = (x["一级模块"], x["二级模块"], x["三级模块"], x["功能过程"])
    groups.setdefault(key, []).append(x)

print("功能过程总数:", len(groups))
print("三级模块总数:", len({k[2] for k in groups}))
print("唯一功能过程名称数:", len({k[3] for k in groups}))
print("\n每个功能过程的子过程数分布:", sorted(Counter(len(v) for v in groups.values()).items()))
print("每个功能过程的子过程类型组合 TOP:",
      Counter(tuple(y["数据移动类型"] for y in v) for v in groups.values()).most_common(8))
print("\n每个三级模块的功能过程数分布:",
      sorted(Counter(sum(1 for k in groups if k[2] == l3) for l3 in {k[2] for k in groups}).items()))

names = sorted({k[3] for k in groups})
head_c, tail_c, other = Counter(), Counter(), []
for n in names:
    h = next((v for v in VERBS if n.startswith(v)), None)
    t = next((v for v in VERBS if n.endswith(v)), None)
    if h:
        head_c[h] += 1
    elif t:
        tail_c[t] += 1
    else:
        other.append(n)
print(f"\n功能过程命名：动词在句首 {sum(head_c.values())} 条 -> {head_c.most_common()}")
print(f"            动词在句尾 {sum(tail_c.values())} 条 -> {tail_c.most_common()}")
print(f"            两端都不是 {len(other)} 条 -> 示例: {other[:8]}")

# where does the verb sit inside 子过程描述
sub_head = Counter()
for x in records:
    s = x["子过程描述"]
    v = next((v for v in VERBS if s.startswith(v)), None)
    sub_head[v or "(其他)"] += 1
print("\n子过程描述动词位置:", sub_head.most_common(8))

# a few full module expansions to copy the granularity style
print("\n===== 粒度样例 =====")
for target in ("话费收取", "预存清退", "IMS详单批处理管控", "综合退费管理批处理管控"):
    fps = [(k, v) for k, v in groups.items() if k[2] == target]
    if not fps:
        continue
    print(f"\n[{target}] 三级模块 / 功能过程 {len(fps)} 个")
    for k, v in fps:
        print(f"   - {k[3]}   子过程: " + " / ".join(
            f"{y['数据移动类型']}:{y['数据组']}" for y in v))
