# -*- coding: utf-8 -*-
"""Independent audit of the filled 功能过程 list: verbs, near-duplicates, 凑数 risk."""
import difflib
import sys
from collections import Counter

import openpyxl

sys.path.insert(0, r"d:\Projects\Persion\ai-customer-service\fun")
sys.stdout.reconfigure(encoding="utf-8")

from fill_function_processes import PLAN, VERBS  # noqa: E402

PATH = r"C:\Users\16680\Desktop\调账相关.xlsx"
ws = openpyxl.load_workbook(PATH)["Sheet1"]

items = {}
for row in sorted(PLAN):
    lines = [x for x in (ws.cell(row=row, column=5).value or "").split("\n") if x.strip()]
    names = [x.split(".", 1)[1] for x in lines]
    assert names == PLAN[row][1], f"E{row} 与计划不一致"
    items[row] = names

all_names = [(r, n) for r, ns in items.items() for n in ns]
print(f"读回 {len(items)} 行 / {len(all_names)} 个功能过程\n")

# 1. verb coverage
heads = Counter(next(v for v in VERBS if n.startswith(v)) for _, n in all_names)
tails = [n for _, n in all_names if n.endswith(tuple(VERBS)) and not n.startswith(tuple(VERBS))]
unused = [v for v in VERBS if v not in heads and v not in "".join(tails)]
print("动词在句首:", sum(heads.values()), dict(heads))
print("动词仅在句尾:", len(tails), tails)
print("未使用到的动词:", unused, "\n")

# 2. near-duplicate names across modules (the real 重复 risk)
print("=== 高相似名称对 (ratio>=0.75，跨行) ===")
flag = 0
for i in range(len(all_names)):
    for j in range(i + 1, len(all_names)):
        (r1, n1), (r2, n2) = all_names[i], all_names[j]
        if r1 == r2:
            continue
        ratio = difflib.SequenceMatcher(None, n1, n2).ratio()
        if ratio >= 0.75:
            flag += 1
            print(f"  {ratio:.2f}  行{r1} {n1}  <->  行{r2} {n2}")
print("(无)" if not flag else f"共 {flag} 对需人工确认")

# 3. 凑数 risk: inside one module, items whose object is a mere step of another
print("\n=== 同模块内对象高度重叠（需确认是否为独立业务对象） ===")
hit = 0
for row, ns in items.items():
    for i in range(len(ns)):
        for j in range(i + 1, len(ns)):
            core = lambda s: s[2:]  # strip leading verb
            a, b = core(ns[i]), core(ns[j])
            if difflib.SequenceMatcher(None, a, b).ratio() >= 0.7:
                hit += 1
                print(f"  行{row}: {ns[i]}  <->  {ns[j]}")
print("(无)" if not hit else f"共 {hit} 对")

# 4. object coverage per module: distinct 业务对象 count should equal item count
print("\n=== 各模块功能过程全文 ===")
for row, ns in items.items():
    print(f"\nE{row}  [{ws.cell(row=row, column=3).value}]  {len(ns)} 个")
    for k, n in enumerate(ns, 1):
        print(f"   {k}.{n}")
