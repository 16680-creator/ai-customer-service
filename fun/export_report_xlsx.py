"""把 COSMIC 拆分成果导出成一份可交付的 Excel 报告。

数据源单一：明细全部回读桌面 调账相关.xlsx，不复制第二份名称清单；
「拆分调整说明」取自 canvas 报告的 MODULES.note，保持两处结论一致。
"""

import datetime
import math
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = r"C:\Users\16680\Desktop\调账相关.xlsx"
OUT = r"C:\Users\16680\Desktop\调账相关-功能过程拆分报告.xlsx"
CANVAS = (r"C:\Users\16680\.qoder\projects\D--Projects-Persion-ai-customer-service"
          r"\canvases\cosmic-split-completion-report.canvas.tsx")

VERBS = ("查询 获取 生成 新增 保存 存管 归档 冻结 申请 提交 修改 更新 删除 汇总 推送 采集 "
         "处理 打印 上传 下载 下发 设置 计算 分析 通知 统计 展示 同步 迁移 预览 审批 对比 "
         "比对 导入 导出").split()

UNIT = 10000 / 916 / 2.5
ROWS = list(range(18, 38))

HEAD_FILL = PatternFill("solid", fgColor="44546A")
BAND_FILL = PatternFill("solid", fgColor="F2F5F9")
TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color="1F3864")
HEAD_FONT = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10.5)
SMALL_FONT = Font(name="微软雅黑", size=10)
WARN_FONT = Font(name="微软雅黑", size=10.5, bold=True, color="C55A11")
EDGE = Side(style="thin", color="BFBFBF")
BOX = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)


def width_of(text):
    """Excel 列宽以数字字符为单位，CJK 字形约占两倍。"""
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in str(text))


def half_up(value, digits=0):
    """Excel 的四舍五入是 half-up，Python 的 round 是银行家舍入，整.5 时会差一档。"""
    factor = 10 ** digits
    return math.floor(value * factor + 0.5) / factor


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def verb_of(name):
    for verb in VERBS:
        if name.startswith(verb):
            return verb, "句首"
    for verb in VERBS:
        if name.endswith(verb):
            return verb, "句尾"
    return "", "未命中"


# --- 读取交付物 ---------------------------------------------------------------
src = load_workbook(SRC)["Sheet1"]
merged = {}
for rng in src.merged_cells.ranges:
    top = src.cell(row=rng.min_row, column=rng.min_col).value
    for row in src[rng.coord]:
        for cell in row:
            merged[(cell.row, cell.column)] = top


def cell(row, col):
    return merged.get((row, col), src.cell(row=row, column=col).value)


notes = dict(re.findall(r'row:\s*"(\d{2})"[\s\S]*?note:\s*"([^"]+)"',
                        open(CANVAS, encoding="utf-8").read()))
modules = {}
for row in ROWS:
    names = [line.split(".", 1)[1] for line in (cell(row, 5) or "").split("\n") if line.strip()]
    modules[row] = {
        "l1": cell(row, 1) or "", "l2": cell(row, 2) or "", "l3": cell(row, 3) or "",
        "effort": cell(row, 4), "owner": cell(row, 7), "names": names,
        "raw": cell(row, 4) * UNIT, "target": round(cell(row, 4) * UNIT, 2),
        "note": notes.get(str(row), ""),
    }

total = sum(len(m["names"]) for m in modules.values())
raw_total = round(sum(m["target"] for m in modules.values()), 3)
assert total == 124 and len(modules) == 20, f"拆分结果已变化：{total}"

wb = Workbook()
today = datetime.date.today().isoformat()

# --- Sheet1 报告概览 ----------------------------------------------------------
ov = wb.active
ov.title = "报告概览"
ov["A1"] = "COSMIC 功能点拆分填写完成报告"
ov["A1"].font = TITLE_FONT
ov.merge_cells("A1:D1")
ov["A2"] = (f"调账相关.xlsx · 小威负责的第 18–37 行共 {len(ROWS)} 个三级模块 · "
            f"粒度基准取自附件3《湖南个人安全-COSMIC功能点拆分表》“个人客户信息安全”表 · {today}")
ov["A2"].font = SMALL_FONT
ov["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ov.merge_cells("A2:D2")
ov.row_dimensions[2].height = 30

facts = [
    ("负责人", "小威"),
    ("覆盖三级模块", f"{len(ROWS)} / {len(ROWS)}（E18:E37 无空行）"),
    ("工作量合计", f"{round(sum(m['effort'] for m in modules.values()), 1)} W"),
    ("F 列换算原值合计", f"{raw_total}（D 合计 {round(sum(m['effort'] for m in modules.values()), 1)} W × 10000÷916÷2.5）"),
    ("取整口径", "先求和再四舍五入 = 124（逐行取整再求和为 126，未采用）"),
    ("实列功能过程", f"{total} 个，与取整目标一致"),
    ("动词命中", "124 / 124 命中指定 35 动词且全部位于句首（覆盖 31 个动词）"),
    ("越出 E18:E37 的改动", "0（合并区、公式、行高、他人行逐格差分确认）"),
]
r = 4
ov.cell(row=r, column=1, value="项目").font = HEAD_FONT
ov.cell(row=r, column=2, value="结论").font = HEAD_FONT
style_header(ov, r, 2)
for k, v in facts:
    r += 1
    ov.cell(row=r, column=1, value=k).font = BODY_FONT
    ov.cell(row=r, column=2, value=v).font = BODY_FONT
    for c in (1, 2):
        cellobj = ov.cell(row=r, column=c)
        cellobj.border = BOX
        cellobj.alignment = Alignment(wrap_text=True, vertical="center")

paras = [
    ("采用的取整口径",
     "F 列公式为 工作量×10000÷916÷2.5，20 行原始值合计 124.454。逐行 round 再求和得 126，"
     "先求和再四舍五入得 124，本次按你确定的 124 执行。削减落在两个向上取整幅度最大（各 +0.45）"
     "且业务上最可压缩的 6.55 行：E18 去掉“保存常用查询条件”、E24 去掉“导出规则清单”。"
     "因此这两行会出现“F 显示 6.6、实列 6 条”。其余 18 行的实列条数等于各自 F 原始值逐行四舍五入的结果；"
     "另 E31 的 F 列显示 3.5，但原始值 3.4934 四舍五入为 3 条——这两类口径差异已在“模块分配对比”表单列。"),
    ("从附件3 提取的判定性规律",
     "按合并区还原后复现出 336 个功能过程、859 条数据移动、118 个三级模块。决定性判据是每个功能过程的"
     "子过程数只落在 2 或 3：E+R+X 共 187 个、E+W 共 145 个、E+X 共 4 个。所以“业务对象或触发事件不同”"
     "才构成两个功能过程，而“输入→读取→返回结果”永远只算一条。附件3 中“查询话费收取账本”与"
     "“查询涉敏客户资料”即为两条，本次拆分沿用该边界。"),
    ("拆分自检与清理",
     "跨模块相似度扫描（ratio≥0.75）定位到 1 处真重复、3 处内部步骤凑数，已全部改名："
     "33.1 与 35.1 的“采集酬金退费/退款明细数据”重复 → 33.1 改为“查询酬金口径退费账目数据”；"
     "22.3“更新权限状态”疑为迁移内部步骤 → “冻结政企调账页面旧权限”；"
     "21.3“保存校验流水” → “设置政企调账规则校验参数”；34.4“保存差异明细” → “统计酬金稽核差异退费金额”。"
     "修正后余 6 对合规同类（规则↔模板、BOSS↔OA、供数↔稽核归档）。"),
]
for title, body in paras:
    r += 2
    ov.cell(row=r, column=1, value=title).font = Font(name="微软雅黑", size=12, bold=True, color="1F3864")
    r += 1
    ov.cell(row=r, column=1, value=body).font = BODY_FONT
    ov.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ov.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ov.row_dimensions[r].height = max(46, 15 * (width_of(body) // 78 + 1))
ov.column_dimensions["A"].width = 26
ov.column_dimensions["B"].width = 96

# --- Sheet2 功能过程明细 ------------------------------------------------------
dt = wb.create_sheet("功能过程明细")
heads = ["行号", "一级模块", "二级模块", "三级模块", "模块功能过程数", "序号", "功能过程名称", "动词", "动词位置"]
dt.append(heads)
style_header(dt, 1, len(heads))
band, row_ptr = False, 1
for row in ROWS:
    m = modules[row]
    band = not band
    for i, name in enumerate(m["names"], 1):
        verb, pos = verb_of(name)
        row_ptr += 1
        vals = [f"E{row}", m["l1"], m["l2"], m["l3"], len(m["names"]), i, name, verb, pos]
        for c, val in enumerate(vals, 1):
            cellobj = dt.cell(row=row_ptr, column=c, value=val)
            cellobj.font = BODY_FONT
            cellobj.border = BOX
            cellobj.alignment = Alignment(
                horizontal="center" if c in (1, 5, 6, 9) else "left",
                vertical="center", wrap_text=c in (2, 3, 4, 7))
            if band:
                cellobj.fill = BAND_FILL
    dt.row_dimensions[row_ptr].height = 18
dt.freeze_panes = "E2"
dt.auto_filter.ref = f"A1:I{row_ptr}"
for letter, w in zip("ABCDEFGHI", (8, 18, 22, 34, 15, 7, 40, 9, 10)):
    dt.column_dimensions[letter].width = w
print(f"明细 {row_ptr - 1} 行数据")

# --- Sheet3 模块分配对比 ------------------------------------------------------
cmp_ws = wb.create_sheet("模块分配对比")
heads = ["单元格", "二级模块", "三级模块", "工作量（W）", "F 列原始值", "F 列显示值",
         "逐行四舍五入", "实列条数", "与逐行取整之差", "数量口径说明", "拆分调整说明"]
cmp_ws.append(heads)
style_header(cmp_ws, 1, len(heads))
for i, row in enumerate(ROWS, 2):
    m = modules[row]
    raw, listed = m["raw"], len(m["names"])
    shown, per_row = half_up(raw, 1), int(half_up(raw))
    if listed != per_row:
        remark = f"逐行四舍五入应为 {per_row} 条，为满足合计 124 记为 {listed} 条"
    elif int(half_up(shown)) != listed:
        remark = f"F 列显示 {shown}，原始值 {raw:.4f} 四舍五入实为 {listed} 条"
    else:
        remark = ""
    vals = [f"E{row}", m["l2"], m["l3"], m["effort"], round(raw, 4), shown, per_row,
            listed, listed - per_row, remark, m["note"]]
    for c, val in enumerate(vals, 1):
        cellobj = cmp_ws.cell(row=i, column=c, value=val)
        cellobj.font = WARN_FONT if (c == 9 and val) else BODY_FONT
        cellobj.border = BOX
        cellobj.alignment = Alignment(
            horizontal="center" if c in (1, 4, 5, 6, 7, 8, 9) else "left",
            vertical="center", wrap_text=c in (3, 10, 11))
        cellobj.number_format = {4: "0.0", 5: "0.0000", 6: "0.0"}.get(c, "General")
tr = len(ROWS) + 2
cmp_ws.cell(row=tr, column=3, value="合计")
for col, fmt in ((4, "0.0"), (5, "0.0000"), (7, "0"), (8, "0"), (9, "0")):
    letter = get_column_letter(col)
    cmp_ws.cell(row=tr, column=col,
                value=f"=SUM({letter}2:{letter}{tr - 1})").number_format = fmt
for c in range(1, len(heads) + 1):
    head = cmp_ws.cell(row=tr, column=c)
    head.fill = HEAD_FILL
    head.font = HEAD_FONT
    head.border = BOX
    head.alignment = Alignment(horizontal="center", vertical="center")
cmp_ws.cell(row=tr + 1, column=3,
           value="合计行说明：逐行取整共 126，实列 124，差 2 条即 E18/E24 各减 1").font = SMALL_FONT
for letter, w in zip("ABCDEFGHIJK", (9, 22, 32, 11, 12, 11, 13, 11, 17, 38, 44)):
    cmp_ws.column_dimensions[letter].width = w

chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "各三级模块：实列功能过程条数 vs F 列换算原值"
chart.y_axis.title = "功能过程（个）"
chart.x_axis.title = "单元格（E18–E37）"
chart.height, chart.width = 11, 30
for col in (5, 8):
    chart.add_data(Reference(cmp_ws, min_col=col, max_col=col, min_row=1, max_row=tr - 1),
                   titles_from_data=True)
chart.set_categories(Reference(cmp_ws, min_col=1, min_row=2, max_row=tr - 1))
cmp_ws.add_chart(chart, "A26")

# --- Sheet4 逐项校验证据 ------------------------------------------------------
ev = wb.create_sheet("逐项校验证据")
ev.append(["计划要求", "校验方式", "结果"])
style_header(ev, 1, 3)
evidence = [
    ("条数＝总数 124", "回读 E18:E37 按换行切分逐格断言等于计划；换行符 104 = 124−20", "通过"),
    ("编号＋名称逐条换行", "断言每格各行严格等于 “1.名称”，无空格、无嵌套编号", "通过"),
    ("动词在句首或句尾", "35 个指定动词双端匹配：句首 124 条、仅在句尾 0 条", "通过"),
    ("检查重复", "跨模块相似度扫描 ratio≥0.75：真重复 1 处已消除，余 6 对为合规同类", "通过"),
    ("不越出模块范围", "18–23 全含“政企”、24–31 全含“审批”、32 全含“经分报表”、33–35 全含“酬金”、36–37 全含“结算”", "通过"),
    ("保留模块与负责人", "全工作簿逐格差分：取值差异 20 处全部落在 E18:E37，越界改动 0", "通过"),
    ("保留数量公式", "F18:F37 与 D59 合计公式逐字符相等；D 工作量、G 负责人共 60 格未变", "通过"),
    ("必要时调整行高", "行高差异仅 20 行且全在 18–37；最宽单行 30/65.1 列宽单位占 46%，行高 条数×17pt ≥ 所需 15pt/行", "通过"),
    ("合并区与样式不回归", "15 个合并区、列宽集合、条件格式、数据有效性、图片、图表、批注、冻结窗格、自动筛选全部一致；"
     "唯一样式改动为 E18:E37 水平对齐 center→left", "通过"),
    ("不误填他人行", "第 2–17、38–58 行 E 列扫描为空", "通过"),
]
for i, (req, how, res) in enumerate(evidence, 2):
    for c, val in enumerate((req, how, res), 1):
        cellobj = ev.cell(row=i, column=c, value=val)
        cellobj.font = BODY_FONT
        cellobj.border = BOX
        cellobj.alignment = Alignment(horizontal="center" if c == 3 else "left",
                                      vertical="center", wrap_text=True)
    ev.cell(row=i, column=3).font = Font(name="微软雅黑", size=10.5, bold=True, color="2E7D32")
    ev.row_dimensions[i].height = 34
for letter, w in zip("ABC", (24, 96, 10)):
    ev.column_dimensions[letter].width = w
ev.freeze_panes = "A2"

# --- Sheet5 产物与待办 --------------------------------------------------------
ar = wb.create_sheet("产物与待办")
ar.append(["路径", "性质", "说明"])
style_header(ar, 1, 3)
artifacts = [
    (SRC, "交付物", "原表，已原位写入 E18:E37 共 124 条与 18–37 行高"),
    (r"C:\Users\16680\Desktop\调账相关.backup-20260905-111154.xlsx", "还原基线",
     "写入前的纯净原件，差分参照，建议保留"),
    (OUT, "本报告", "由拆分结果生成的多 Sheet 报告，可独立转发"),
    ("fun/fill_function_processes.py", "生成脚本", "PLAN＋静态校验＋等待解锁＋原位写入＋回读审计，改任一条重跑即可"),
    ("fun/analyze_cosmic_ref.py", "校验脚本", "解析附件3，按合并区还原后统计粒度与命名规律"),
    ("fun/audit_split.py", "校验脚本", "动词位置、跨模块近似重复、同模块凑数风险扫描"),
    ("fun/diff_workbooks.py", "校验脚本", "原件与纯净基线的全工作簿逐格差分，含工作簿级对象回归检查"),
    ("fun/verify_filled.py", "校验脚本", "逐行回读条数、合计、空行与他人行误填检查"),
    ("fun/check_display_fit.py", "校验脚本", "显示度量：单行是否二次折行、行高是否足够不裁切"),
    ("fun/verify_canvas_report.py", "校验脚本", "报告数字与 xlsx 单元格逐字一致性核对"),
]
for i, (path, kind, desc) in enumerate(artifacts, 2):
    for c, val in enumerate((path, kind, desc), 1):
        cellobj = ar.cell(row=i, column=c, value=val)
        cellobj.font = SMALL_FONT
        cellobj.border = BOX
        cellobj.alignment = Alignment(horizontal="center" if c == 2 else "left",
                                      vertical="center", wrap_text=True)
    ar.row_dimensions[i].height = 30

pend = [("桌面 3 份冗余文件请手动删除",
         "调账相关-已填写.xlsx、调账相关.backup-20260905-111738.xlsx、"
         "调账相关.backup-20260905-111829.xlsx 均为过期中间产物；我的工具不允许删除项目目录外的文件"),
        ("E18:E37 水平对齐是否改回居中",
         "现为 left，因为多行列表居中会错位；想保持居中说一句即可改回"),
        ("E30 第 3 条措辞是否还原",
         "现为“比对客服审批页面嵌入来源域名”，替换了更早的“展示嵌入的客服审批页面”——后者更像输出步骤")]
start = len(artifacts) + 3
ar.cell(row=start, column=1, value="需要你确认或手动处理").font = Font(
    name="微软雅黑", size=12, bold=True, color="1F3864")
for i, (title, detail) in enumerate(pend, start + 1):
    ar.cell(row=i, column=1, value=f"{i - start}. {title}").font = BODY_FONT
    ar.cell(row=i, column=2, value=detail).font = SMALL_FONT
    ar.cell(row=i, column=3, value="待你决定").font = WARN_FONT
    for c in range(1, 4):
        ar.cell(row=i, column=c).border = BOX
        ar.cell(row=i, column=c).alignment = Alignment(
            horizontal="center" if c == 3 else "left", vertical="center", wrap_text=True)
    ar.row_dimensions[i].height = 34
for letter, w in zip("ABC", (44, 12, 78)):
    ar.column_dimensions[letter].width = w

# openpyxl 无缓存值，公式所在表保持计算模式，交由 Excel 打开时求值
wb.calculation.fullCalcOnLoad = True

try:
    wb.save(OUT)
except PermissionError:
    stamp = datetime.datetime.now().strftime("%H%M%S")
    OUT = OUT.replace(".xlsx", f"-{stamp}.xlsx")
    wb.save(OUT)
print(f"已写出：{OUT}")
print(f"Sheet：{wb.sheetnames}")
sys.exit(0)
