# -*- coding: utf-8 -*-
"""Fill 功能过程 (E18:E37) of 调账相关.xlsx with COSMIC-style splits.

Conventions taken from 附件3《湖南个人安全-COSMIC功能点拆分表》:
  * one 功能过程 = one business object / one trigger event, carrying 2~3 子过程
    (E+R+X for query/export style, E+W for create/update style);
  * internal 输入/读取/写入/输出 steps are NOT counted as separate 功能过程;
  * name = 动词 + 业务对象 (+限定语), verb placed at head or tail.

Allocation target: 124 功能过程 = round(sum(F)) i.e. 合计取整, not 逐行取整(126).
"""
import shutil
import sys
import time
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment

sys.stdout.reconfigure(encoding="utf-8")

SRC = r"C:\Users\16680\Desktop\调账相关.xlsx"
UNIT = 10000 / 916 / 2.5  # 功能过程 per person-week, same factor as column F

VERBS = """查询 获取 生成 新增 保存 存管 归档 冻结 申请 提交 修改 更新 删除 汇总 推送 采集 处理
打印 上传 下载 下发 设置 计算 分析 通知 统计 展示 同步 迁移 预览 审批 对比 比对 导入 导出""".split()

# row -> (三级模块 must-contain keyword, [功能过程 names])
PLAN = {
    18: ("政企", [
        "查询政企客户基本档案信息",
        "查询政企客户账户余额信息",
        "查询政企客户历史调账记录",
        "查询政企可调账费用明细数据",
        "获取政企调账页面初始化要素",
        "新增政企调账页面操作埋点记录",
    ]),
    19: ("政企", [
        "查询政企集团客户产品信息",
        "查询政企成员号码产品信息",
        "查询政企产品费用账单信息",
        "采集政企产品查询操作埋点数据",
        "导出政企客户产品查询结果",
    ]),
    20: ("政企", [
        "新增政企调账单基本信息",
        "新增政企调账单费用明细",
        "修改政企调账单信息",
        "删除政企调账单草稿",
        "提交政企调账申请工单",
    ]),
    21: ("政企", [
        "比对政企调账金额限额规则",
        "比对政企调账频次超限规则",
        "设置政企调账规则校验参数",
        "查询政企调账规则校验结果",
    ]),
    22: ("政企", [
        "查询政企调账页面权限配置",
        "迁移政企调账页面操作权限",
        "冻结政企调账页面旧权限",
        "归档政企调账页面权限迁移记录",
    ]),
    23: ("政企", [
        "查询个人政企界面隔离标识",
        "设置个人政企界面隔离规则",
        "处理个人政企界面越权访问请求",
    ]),
    24: ("审批", [
        "新增退费审批规则",
        "修改退费审批规则",
        "删除退费审批规则",
        "查询退费审批规则配置",
        "设置退费审批规则生效范围",
        "提交退费审批规则变更申请",
    ]),
    25: ("审批", [
        "查询退费金额审批权限层级",
        "计算退费次数审批阈值",
        "生成退费审批节点流转路线",
        "审批退费工单节点任务",
        "更新退费审批节点处理状态",
        "推送退费审批节点待办通知",
    ]),
    26: ("审批", [
        "新增退费审批模板",
        "修改退费审批模板",
        "删除退费审批模板",
        "查询退费审批模板配置",
        "预览退费审批单模板效果",
        "打印退费审批单据",
        "下载退费审批附件材料",
    ]),
    27: ("审批", [
        "提交退费工单至BOSS审批流程",
        "获取BOSS审批流程受理结果",
        "查询BOSS审批流程工单状态",
        "更新BOSS审批流程回写状态",
        "同步BOSS审批流程节点信息",
        "处理BOSS审批流程异常回执",
        "推送BOSS审批流程催办通知",
        "归档BOSS审批流程结束工单",
    ]),
    28: ("审批", [
        "提交退费工单至OA审批流程",
        "获取OA审批流程待办任务",
        "查询OA审批流程实例状态",
        "更新OA审批流程回写结果",
        "同步OA审批组织部门信息",
        "对比OA与BOSS审批状态一致性",
        "处理OA审批流程异常报文",
        "生成OA审批联调测试数据",
        "归档OA审批流程结束实例",
    ]),
    29: ("审批", [
        "新增BOSS审批工号",
        "新增OA审批工号",
        "修改审批工号信息",
        "删除审批工号信息",
        "查询BOSS与OA审批工号关系",
        "设置BOSS与OA审批工号映射关系",
        "同步审批工号组织架构信息",
        "导入审批工号关系配置数据",
        "导出审批工号关系配置清单",
    ]),
    30: ("审批", [
        "查询客服审批页面嵌入配置",
        "获取客服审批页面免登跳转地址",
        "比对客服审批页面嵌入来源域名",
        "更新客服审批页面嵌入状态",
    ]),
    31: ("审批", [
        "查询退费审批单当前状态",
        "查询退费审批流转轨迹明细",
        "导出退费审批状态轨迹清单",
    ]),
    32: ("经分报表", [
        "采集经分报表退费扣回数据",
        "汇总经分报表退费指标数据",
        "计算经分报表退费分摊金额",
        "更新华为经分报表取数口径",
        "生成华为经分报表上报文件",
        "上传华为经分报表结果文件",
        "查询华为经分报表加工结果",
        "导出华为经分报表明细清单",
        "归档华为经分报表历史账期数据",
    ]),
    33: ("酬金", [
        "查询酬金口径退费账目数据",
        "汇总酬金扣回账期数据",
        "计算酬金分摊退费金额",
        "生成酬金供数接口文件",
        "下发酬金供数接口文件数据",
        "查询酬金供数任务执行状态",
        "更新酬金供数文件落地状态",
        "推送酬金供数失败补发数据",
        "归档酬金供数历史文件",
    ]),
    34: ("酬金", [
        "查询酬金稽核文件清单",
        "下载酬金稽核文件",
        "处理酬金稽核文件数据",
        "统计酬金稽核差异退费金额",
        "更新酬金稽核文件处理状态",
        "通知酬金稽核差异处理结果",
        "归档酬金稽核历史文件",
    ]),
    35: ("酬金", [
        "采集酬金业务退款明细数据",
        "导入酬金退款对接数据文件",
        "导出酬金退款待对接数据清单",
        "汇总酬金退款账期对账数据",
        "比对酬金退款双方差异数据",
        "处理酬金退款差异调整数据",
        "更新酬金退款数据对接状态",
        "推送酬金退款对接结果反馈",
        "生成酬金退款对账差异报告",
    ]),
    36: ("结算", [
        "查询结算数据文件清单",
        "下载渠道结算数据文件",
        "采集结算退费归集数据",
        "汇总结算数据账期归集结果",
        "保存结算数据归集明细",
        "统计结算退费归集金额",
        "归档结算数据历史文件",
    ]),
    37: ("结算", [
        "查询结算业务退费明细数据",
        "更新结算业务退费处理规则",
        "计算结算业务退费扣回金额",
        "处理结算业务退费冲正数据",
    ]),
}

LINE_H = 17.0  # pt per wrapped line for 微软雅黑 11 (sheet base is 16.5)
POLL = 5.0     # seconds between write-lock retries


def is_locked(path):
    """WPS/Excel holds an exclusive handle on an open workbook, so probe with r+b."""
    try:
        open(path, "r+b").close()
        return False
    except OSError:
        return True


def wait_for_unlock(seconds):
    """Poll until the original file is writable, so the in-place write lands the moment
    the user closes the workbook. Returns False on timeout."""
    deadline = time.time() + seconds
    while is_locked(SRC):
        left = int(deadline - time.time())
        if left <= 0:
            return False
        print(f"原文件仍被 WPS/Excel 占用，{POLL:.0f}s 后重试（剩余 {left}s）...", flush=True)
        time.sleep(POLL)
    return True


def validate():
    """Static checks before touching the workbook."""
    errs = []
    total = sum(len(v[1]) for v in PLAN.values())
    if total != 124:
        errs.append(f"总数 {total} != 124")

    seen = {}
    for row, (scope, names) in sorted(PLAN.items()):
        if not (18 <= row <= 37):
            errs.append(f"行 {row} 越出 18-37")
        for i, n in enumerate(names, 1):
            if not n.startswith(tuple(VERBS)) and not n.endswith(tuple(VERBS)):
                errs.append(f"E{row} 第{i}条动词不在首尾: {n}")
            if scope not in n:
                errs.append(f"E{row} 第{i}条越出模块范围(缺'{scope}'): {n}")
            if len(n) > 20:
                errs.append(f"E{row} 第{i}条过长({len(n)}): {n}")
            if n in seen:
                errs.append(f"重复名称: {n} (行{seen[n]} 与 行{row})")
            seen[n] = row
            if n.startswith(f"{i}.") or "\n" in n:
                errs.append(f"E{row} 第{i}条自带编号或换行: {n}")
    return total, errs


def main():
    total, errs = validate()
    if errs:
        print("校验未通过，未写入文件：")
        for e in errs:
            print("  !", e)
        sys.exit(1)
    print(f"静态校验通过：{len(PLAN)} 个三级模块 / {total} 个功能过程\n")

    # WPS/Excel keeps an exclusive handle on the open workbook; probe before touching it
    if "--wait" in sys.argv:
        budget = int(sys.argv[sys.argv.index("--wait") + 1])
        if not wait_for_unlock(budget):
            print(f"等待 {budget}s 超时，原文件仍被占用，本次未写入。")
            sys.exit(2)
        print("原文件已释放句柄，开始原位写入。\n", flush=True)

    target = SRC
    backup = None
    if is_locked(SRC):
        target = SRC.replace(".xlsx", "-已填写.xlsx")
        print(f"原文件不可写（WPS/Excel 占用中），降级为写出已填写副本 -> {target}\n")
    else:
        backup = SRC.replace(".xlsx", f".backup-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
        shutil.copy2(SRC, backup)
        print(f"已备份原文件 -> {backup}", flush=True)

    wb = openpyxl.load_workbook(SRC)
    ws = wb["Sheet1"]

    # snapshot every other column of the 20 rows so we can prove nothing else moved
    before = {(r, c): ws.cell(row=r, column=c).value
              for r in PLAN for c in (1, 2, 3, 4, 6, 7, 8)}

    for row, (_scope, names) in sorted(PLAN.items()):
        cell = ws.cell(row=row, column=5)
        cell.value = "\n".join(f"{i}.{name}" for i, name in enumerate(names, 1))
        al = cell.alignment
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True,
                                   indent=al.indent)
        ws.row_dimensions[row].height = round(len(names) * LINE_H, 1)

    # the 20 rows must keep C 三级模块 / D 工作量 / F 公式 / G 负责人 untouched
    for (r, c), old in before.items():
        new = ws.cell(row=r, column=c).value
        assert new == old, f"{ws.cell(row=r, column=c).coordinate} 被意外改动: {old!r} -> {new!r}"
    for r in PLAN:
        f = ws.cell(row=r, column=6).value
        assert isinstance(f, str) and f.startswith("=D"), f"F{r} 公式被破坏: {f!r}"

    wb.save(target)
    print(f"已写入 {target}")
    report(target, backup)


def report(path, backup):
    """Re-open what we just wrote and audit it. Formulas have no cached value after an
    openpyxl save, so derive F from D x UNIT and assert the formula text separately."""
    wb_f = openpyxl.load_workbook(path, data_only=False)
    ws = wb_f["Sheet1"]

    print("\n行  三级模块                                工作量  F原值   条数  行高")
    print("-" * 84)
    acc = raw_sum = 0
    for row in sorted(PLAN):
        names = PLAN[row][1]
        cell = ws.cell(row=row, column=5).value or ""
        lines = [x for x in cell.split("\n") if x.strip()]
        d = ws.cell(row=row, column=4).value
        formula = ws.cell(row=row, column=6).value
        h = ws.row_dimensions[row].height
        l3 = ws.cell(row=row, column=3).value
        assert lines == [f"{i}.{n}" for i, n in enumerate(names, 1)], f"E{row} 落盘内容不符"
        assert formula == f"=D{row}*10000/916/2.5", f"F{row} 公式意外: {formula!r}"
        assert ws.cell(row=row, column=7).value == "小威", f"G{row} 负责人被改动"
        fv = d * UNIT
        acc += len(lines)
        raw_sum += fv
        print(f"{row:<3} {l3:<34} {d:<6} {fv:<7.4f} {len(lines):<4} {h}")
    print("-" * 84)
    print(f"合计 功能过程 {acc} 个   F列原始值合计 {raw_sum:.3f}   取整目标 {round(raw_sum)}")
    assert acc == 124, acc
    print("\n校验：条数一致 / 编号 1. 格式 / F 列公式与 D、G 列未改动")
    if backup:
        print(f"备份文件：{backup}")
    else:
        print("注意：本次写的是副本，原文件未改动。")


if __name__ == "__main__":
    main()
