# -*- coding: utf-8 -*-
"""Dump structure of the 调账 workbook so we can plan the E-column fill."""
import sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

SRC = r"C:\Users\16680\Desktop\调账相关.xlsx"

wb_f = openpyxl.load_workbook(SRC, data_only=False)
wb_v = openpyxl.load_workbook(SRC, data_only=True)

print("sheets:", wb_f.sheetnames)

for name in wb_f.sheetnames:
    ws_f = wb_f[name]
    ws_v = wb_v[name]
    print("=" * 70)
    print(f"[{name}] dims={ws_f.dimensions} max_row={ws_f.max_row} max_col={ws_f.max_column}")
    print("merged:", sorted(str(r) for r in ws_f.merged_cells.ranges))
    for row in ws_f.iter_rows(min_row=1, max_row=ws_f.max_row, max_col=ws_f.max_column):
        for c in row:
            if c.value is None:
                continue
            v = ws_v.cell(row=c.row, column=c.column).value
            loc = ws_f.merged_cells.ranges
            inmerge = any(c.coordinate in r for r in loc)
            print(f"  {c.coordinate}{'(M)' if inmerge else '    '} f={c.value!r} v={v!r}")
    print("--- row heights ---")
    for r, dim in sorted(ws_f.row_dimensions.items()):
        if dim.height is not None:
            print(f"  row {r}: height={dim.height}")
    print("--- col widths ---")
    for k, dim in sorted(ws_f.column_dimensions.items()):
        if dim.width is not None:
            print(f"  col {k}: width={dim.width}")
