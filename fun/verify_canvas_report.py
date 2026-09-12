"""Check the canvas report against the workbook it claims to describe.

The report is hand-written JSX, so every number in it can silently drift away
from the real E18:E37 content. This re-reads both sides and asserts they agree.
"""

import re
import sys

from openpyxl import load_workbook

XLSX = r"C:\Users\16680\Desktop\调账相关.xlsx"
CANVAS = (
    r"C:\Users\16680\.qoder\projects"
    r"\D--Projects-Persion-ai-customer-service\canvases"
    r"\cosmic-split-completion-report.canvas.tsx"
)
UNIT = 10000 / 916 / 2.5
ROWS = list(range(18, 38))

source = open(CANVAS, encoding="utf-8").read()
failures = []


def check(ok, label, detail=""):
    print(f"[{'OK' if ok else 'FAIL'}] {label}{(' — ' + detail) if detail else ''}")
    if not ok:
        failures.append(label)


# --- workbook side -----------------------------------------------------------
wb = load_workbook(XLSX)
ws = wb["Sheet1"]
actual = {}
for row in ROWS:
    raw = ws.cell(row=row, column=5).value or ""
    actual[row] = [line.split(".", 1)[1] for line in raw.split("\n") if line.strip()]

# --- report side -------------------------------------------------------------
block = source.split("const MODULES: ModuleRow[] = [", 1)[1].split("\n];", 1)[0]
reported = {}
for m in re.finditer(r'\{\s*row:\s*"(\d+)".*?count:\s*(\d+).*?items:\s*\[(.*?)\]', block, re.S):
    row, count, body = int(m.group(1)), int(m.group(2)), m.group(3)
    reported[row] = re.findall(r'"([^"]+)"', body)
    if len(reported[row]) != count:
        check(False, f"row {row}: count field", f"count={count} but {len(reported[row])} items")

check(set(reported) == set(ROWS), "report covers exactly rows 18-37", f"{len(reported)} rows")

mismatch = [r for r in ROWS if reported.get(r) != actual.get(r)]
check(not mismatch, "every listed process matches the cell text", f"diff rows: {mismatch}")

for row in ROWS:
    a, b = actual[row], reported.get(row, [])
    for x, y in zip(a, b):
        if x != y:
            print(f"    row {row}: xlsx={x!r} canvas={y!r}")
            break

counts = [len(v) for v in reported.values()]
check(sum(counts) == 124, "report items total 124", f"sum={sum(counts)}")
check(sum(len(v) for v in actual.values()) == 124, "workbook cells total 124")

const_counts = [int(x) for x in re.findall(r"\d+", source.split("const COUNTS = [", 1)[1].split("];", 1)[0])]
check(const_counts == counts, "COUNTS array matches per-row item counts", f"{const_counts}")
check(len(const_counts) == 20 and sum(const_counts) == 124, "COUNTS has 20 entries summing to 124")

labels = re.findall(r'"(\d{2}) [^"]+"', source.split("const CATEGORY_LABELS = [", 1)[1].split("];", 1)[0])
check([int(x) for x in labels] == ROWS, "CATEGORY_LABELS are rows 18-37 in order")

targets = [
    float(x)
    for x in re.findall(r"[\d.]+", source.split("const TARGETS = [", 1)[1].split("];", 1)[0])
]
check(len(targets) == 20, "TARGETS has 20 entries")
real_targets = [round(ws.cell(row=r, column=4).value * UNIT, 2) for r in ROWS]
bad = [(r, t, x) for r, t, x in zip(ROWS, targets, real_targets) if abs(t - x) > 0.005]
check(not bad, "TARGETS equal D * 10000/916/2.5 (2dp)", f"off: {bad}")

owners = {row: ws.cell(row=row, column=7).value for row in ROWS}
check(set(owners.values()) == {"小威"}, "G18:G37 owner is 小威", str(set(owners.values())))
check(round(sum(ws.cell(row=r, column=4).value for r in ROWS), 2) == 28.5, "D18:D37 sums to 28.5W")

# --- SDK surface -------------------------------------------------------------
sdk = open(r"C:\Users\16680\.qoder\canvas\sdk\index.d.ts", encoding="utf-8").read()
used = re.findall(r"^\s{2}([A-Z]\w+),$", source.split("from \"qoder/canvas\"", 1)[0], re.M)
missing = [name for name in used if not re.search(rf"\b{name}\b", sdk)]
check(not missing, "all imported components exist in the SDK", f"missing: {missing}")

chart_tones = set(re.findall(r'"([^"]+)"', open(
    r"C:\Users\16680\.qoder\canvas\sdk\chart-primitives.d.ts", encoding="utf-8"
).read().split("export type ChartTone = ", 1)[1].split(";", 1)[0]))
chart_series_tones = re.findall(r'\{ name: "[^"]+", data: \w+, tone: "([^"]+)" \}', source)
check(chart_series_tones and all(t in chart_tones for t in chart_series_tones),
      "chart series tones are valid ChartTone", f"used: {chart_series_tones} allowed: {sorted(chart_tones)}")

print(f"\n功能过程：workbook {sum(len(v) for v in actual.values())} / report {sum(counts)}")
sys.exit(1 if failures else 0)
