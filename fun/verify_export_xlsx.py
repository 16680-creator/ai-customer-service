"""回读导出的 Excel 报告，确认它与桌面原表逐项一致。"""

import sys
from collections import defaultdict

from openpyxl import load_workbook

SRC = r"C:\Users\16680\Desktop\调账相关.xlsx"
OUT = r"C:\Users\16680\Desktop\调账相关-功能过程拆分报告.xlsx"
ROWS = list(range(18, 38))
SHEETS = ["报告概览", "功能过程明细", "模块分配对比", "逐项校验证据", "产物与待办"]

failures = []


def check(ok, label, detail=""):
    print(f"[{'OK' if ok else 'FAIL'}] {label}{(' — ' + detail) if detail else ''}")
    if not ok:
        failures.append(label)


live = load_workbook(SRC)["Sheet1"]
expected = {}
for row in ROWS:
    raw = live.cell(row=row, column=5).value or ""
    expected[f"E{row}"] = [line.split(".", 1)[1] for line in raw.split("\n") if line.strip()]

wb = load_workbook(OUT)
check(wb.sheetnames == SHEETS, "Sheet 结构与顺序", str(wb.sheetnames))

dt = wb["功能过程明细"]
heads = [c.value for c in dt[1]]
check(heads == ["行号", "一级模块", "二级模块", "三级模块", "模块功能过程数", "序号",
                "功能过程名称", "动词", "动词位置"], "明细表头 9 列", str(heads))
got = defaultdict(list)
bad_verb, bad_seq, bad_band = [], [], []
for r in range(2, dt.max_row + 1):
    tag, count, seq, name, verb, pos = (dt.cell(row=r, column=c).value
                                        for c in (1, 5, 6, 7, 8, 9))
    got[tag].append(name)
    if pos == "未命中" or not verb:
        bad_verb.append((r, name))
    if seq != len(got[tag]):
        bad_seq.append((r, tag, seq))
    if count != len(expected.get(tag, [])):
        bad_band.append((r, tag, count))

check(dt.max_row - 1 == 124, "明细数据行数 = 124", f"{dt.max_row - 1}")
check(sum(len(v) for v in got.values()) == 124, "按行号聚合后总数 = 124")
check(list(got) == [f"E{r}" for r in ROWS], "行号覆盖 E18…E37 且顺序正确", f"{len(got)} 组")
check(dict(got) == expected, "每条名称与桌面原表单元格逐字一致")
check(not bad_verb, "动词全部命中且在句首/句尾", f"{bad_verb[:3]}")
check(not bad_seq, "组内序号连续 1..n", f"{bad_seq[:3]}")
check(not bad_band, "模块功能过程数与原表一致", f"{bad_band[:3]}")
check(dt.freeze_panes == "E2" and dt.auto_filter.ref == "A1:I125",
      "明细已冻结前 4 列并启用筛选", f"{dt.freeze_panes} / {dt.auto_filter.ref}")

cmp_ws = wb["模块分配对比"]
check([c.value for c in cmp_ws[1]] == ["单元格", "二级模块", "三级模块", "工作量（W）", "F 列原始值",
                                       "F 列显示值", "逐行四舍五入", "实列条数",
                                       "与逐行取整之差", "数量口径说明", "拆分调整说明"],
      "对比表头 11 列", str([c.value for c in cmp_ws[1]][:5]))
counts, per_rows, raws, flagged, remarked = [], [], [], [], []
for r in range(2, 22):
    raws.append(cmp_ws.cell(row=r, column=5).value)
    per_rows.append(cmp_ws.cell(row=r, column=7).value)
    counts.append(cmp_ws.cell(row=r, column=8).value)
    diff = cmp_ws.cell(row=r, column=9).value
    if diff:
        flagged.append((cmp_ws.cell(row=r, column=1).value, diff))
    if cmp_ws.cell(row=r, column=10).value:
        remarked.append(cmp_ws.cell(row=r, column=1).value)
check(counts == [len(expected[f"E{r}"]) for r in ROWS], "对比表实列条数逐行等于明细聚合值")
check(sum(counts) == 124, "对比表实列合计 124", f"{sum(counts)}")
check(sum(per_rows) == 126, "对比表逐行四舍五入合计 126（未被采用的口径）", f"{sum(per_rows)}")
check(abs(sum(raws) - 124.454) < 0.01, "对比表 F 原始值合计 ≈ 124.454", f"{round(sum(raws), 3)}")
check(flagged == [("E18", -1), ("E24", -1)], "仅 E18/E24 比逐行取整少 1 条", str(flagged))
check(remarked == ["E18", "E24", "E31"], "口径说明仅标出这 3 行", str(remarked))
check("3.4934" in str(cmp_ws.cell(row=15, column=10).value), "E31 的显示 3.5 已解释到原始值层面",
      str(cmp_ws.cell(row=15, column=10).value))
tr = 22
check(str(cmp_ws.cell(row=tr, column=8).value) == "=SUM(H2:H21)", "合计行为活公式",
      str(cmp_ws.cell(row=tr, column=8).value))
check(len(cmp_ws._charts) == 1, "对比表内嵌柱状图 1 张", f"{len(cmp_ws._charts)}")

ev = wb["逐项校验证据"]
check(ev.max_row == 11 and all(ev.cell(row=r, column=3).value == "通过" for r in range(2, 12)),
      "校验证据 10 条且结果均为通过")

ar = wb["产物与待办"]
paths = [ar.cell(row=r, column=1).value for r in range(2, 12)]
check(len(paths) == 10 and all(paths), "产物清单 10 项")
check(ar.cell(row=13, column=1).value == "需要你确认或手动处理", "待办区块标题就位")
pend = [ar.cell(row=r, column=1).value for r in range(14, 17)]
check(len(pend) == 3 and all(p and p[0] in "123" for p in pend), "待办 3 项", str([p[:2] for p in pend]))

ov = wb["报告概览"]
body = "\n".join(str(ov.cell(row=r, column=c).value or "") for r in range(1, ov.max_row + 1) for c in (1, 2))
for token in ("124", "124.454", "336", "859", "E+R+X", "政企", "冻结政企调账页面旧权限"):
    check(token in body, f"概览含关键结论「{token}」")

print(f"\n明细 {dt.max_row - 1} 行 / 合计 {sum(counts)} 个功能过程 / 5 个 Sheet")
sys.exit(1 if failures else 0)
