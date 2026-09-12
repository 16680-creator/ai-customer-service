# -*- coding: utf-8 -*-
"""Final read-back of the filled workbook: full 功能过程 list + totals + layout checks."""
import sys
from collections import Counter

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

PATH = r"C:\Users\16680\Desktop\调账相关.xlsx"
UNIT = 10000 / 916 / 2.5
ws = openpyxl.load_workbook(PATH)["Sheet1"]

total = 0
by_l2 = Counter()
L2 = {18: "政企产品专属调账退费", 24: "BOSS与OA分级审批", 32: "经分报表取数优化",
      33: "退费扣回校验闭环"}


def l2_of(row):
    for start, name in sorted(L2.items(), reverse=True):
        if row >= start:
            return name
    return "?"

for row in range(18, 38):
    l3 = ws.cell(row=row, column=3).value
    d = ws.cell(row=row, column=4).value
    raw = d * UNIT
    lines = [x for x in (ws.cell(row=row, column=5).value or "").split("\n") if x.strip()]
    h = ws.row_dimensions[row].height
    longest = max(len(x.split(".", 1)[1]) for x in lines)
    total += len(lines)
    by_l2[l2_of(row)] += len(lines)
    print(f"\nE{row} | {l3} | 工作量{d}W  目标{raw:.2f}  实列{len(lines)}  行高{h}  最长{longest}字")
    for x in lines:
        print(f"      {x}")

print("\n" + "=" * 60)
print(f"18-37 行功能过程合计：{total}")
for name, n in by_l2.most_common():
    print(f"   {name}: {n}")
print(f"D 列合计 {sum(ws.cell(row=r, column=4).value for r in range(18, 38))} W"
      f" -> F 原始合计 {sum(ws.cell(row=r, column=4).value * UNIT for r in range(18, 38)):.3f}")
print("E 列空白行检查:",
      [r for r in range(18, 38) if not (ws.cell(row=r, column=5).value or "").strip()] or "无")
print("其他人员行(2-17,38-58)E列是否被误填:",
      [r for r in list(range(2, 18)) + list(range(38, 59))
       if (ws.cell(row=r, column=5).value or "").strip()] or "未误填")
